"""
시트 한 장을 **격자 모델**(SheetGrid)로 읽는다. xlsx(openpyxl)와 xls(xlrd)가 같은 모델을 만들고,
그 뒤의 표 감지·렌더링(sheet_layout)은 형식을 모른다.

예전 코드의 문제와 여기서 지키는 것:
- 앞 1,000행·100열만 읽고 나머지를 버렸다 → 사용 범위를 끝까지 읽는다(행·열 상한 없음).
- 셀의 원시 값(0.153, datetime)을 그대로 썼다 → 표시 서식을 적용한 글자를 쓴다(cell_format).
- 계산값 없이 저장된 수식 셀은 빈칸이 됐다 → 시트 XML 에서 수식을 찾아 ``=SUM(B2:B4)`` 로 남긴다.
- 큰 파일도 openpyxl 전체 모드로 열어 메모리를 크게 썼다 → 큰 파일은 read_only 스트리밍으로 읽고,
  스트리밍 모드에 없는 병합·숨김 정보는 시트 XML 에서 직접 읽는다.

좌표는 모두 1-based (row, col).
"""

from __future__ import annotations

import datetime as _dt
import html
import io
import logging
import re
import zipfile
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Set, Tuple

from xgen_doc2chunk.core.processor.excel_helper.cell_format import format_cell_value

logger = logging.getLogger("document-processor")

#: 시트 XML 합계가 이보다 크면 openpyxl read_only 로 읽는다(전체 모드는 XML 크기의 약 12배 메모리).
STREAMING_THRESHOLD_BYTES = 20 * 1024 * 1024

#: 표 구조 판단용으로 빈 셀의 테두리를 볼 최대 셀 수. 넘으면 테두리는 보지 않는다(값만으로 판단).
BORDER_SCAN_MAX_CELLS = 400_000

_WS = re.compile(r"\s+")


@dataclass(slots=True)
class GridCell:
    """셀 하나. text 는 엑셀이 보여 주는 글자, kind 는 값의 종류. (셀이 수백만 개일 수 있어 slots)"""

    text: str
    kind: str  # text | number | date | bool | error | formula


@dataclass
class SheetGrid:
    name: str
    cells: Dict[Tuple[int, int], GridCell] = field(default_factory=dict)
    bordered: Set[Tuple[int, int]] = field(default_factory=set)
    merges: List[Tuple[int, int, int, int]] = field(default_factory=list)  # (r1, c1, r2, c2)
    hidden_rows: Set[int] = field(default_factory=set)
    hidden_cols: Set[int] = field(default_factory=set)
    hidden: bool = False  # 숨긴 시트

    def drop_hidden_cells(self) -> None:
        """숨긴 행·열의 셀을 뺀다(옵션)."""
        if not self.hidden_rows and not self.hidden_cols:
            return
        hr, hc = self.hidden_rows, self.hidden_cols
        self.cells = {k: v for k, v in self.cells.items() if k[0] not in hr and k[1] not in hc}
        self.bordered = {k for k in self.bordered if k[0] not in hr and k[1] not in hc}


def _clean_text(text: str) -> str:
    """셀 글자 정리: 앞뒤 공백 제거, 줄바꿈·연속 공백은 한 칸."""
    return _WS.sub(" ", text).strip()


# ──────────────────────────────────────────────────────────────────────────────
# xlsx 계열: 시트 XML 직접 읽기 (병합·숨김·계산값 없는 수식)
# 시트 XML 은 수십 MB 가 될 수 있어 DOM 을 만들지 않고 정규식으로 바이트를 훑는다.
# ──────────────────────────────────────────────────────────────────────────────

_P = rb"(?:[A-Za-z0-9_]+:)?"  # 네임스페이스 접두(대부분의 파일은 없다)
_TAG_SHEET = re.compile(rb"<" + _P + rb"sheet\b([^>]*)/?>", re.S)
_TAG_REL = re.compile(rb"<" + _P + rb"Relationship\b([^>]*)/?>", re.S)
_TAG_ROW = re.compile(rb"<" + _P + rb"row\b([^>]*)>")
_TAG_COL = re.compile(rb"<" + _P + rb"col\b([^>]*)/?>")
_TAG_MERGE = re.compile(rb"<" + _P + rb"mergeCell\b([^>]*)/?>")
_CELL = re.compile(rb"<(" + _P + rb")c\b([^>]*?)(?<!/)>(.*?)</\1c>", re.S)
_F = re.compile(rb"<" + _P + rb"f\b([^>]*?)(?:/>|>(.*?)</" + _P + rb"f>)", re.S)
_V = re.compile(rb"<" + _P + rb"v\b[^>]*>(.*?)</" + _P + rb"v>", re.S)
_ATTR = re.compile(rb"([A-Za-z_][A-Za-z0-9_.:-]*)=\"([^\"]*)\"")
_REF = re.compile(r"([A-Z]+)([0-9]+)")


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def _parse_ref(ref: str) -> Tuple[int, int]:
    m = _REF.fullmatch(ref.replace("$", ""))
    if not m:
        raise ValueError(ref)
    return int(m.group(2)), _col_index(m.group(1))


def _attrs(blob: bytes) -> Dict[str, str]:
    """속성 이름은 접두(r:)를 떼고 돌려준다. 단 r:id 는 'rid' 로 구분한다."""
    out: Dict[str, str] = {}
    for k, v in _ATTR.findall(blob):
        key = k.decode()
        if ":" in key:
            pre, local = key.split(":", 1)
            key = "rid" if local == "id" and pre != "xmlns" else local
        out[key] = html.unescape(v.decode("utf-8", "replace"))
    return out


@dataclass
class SheetXmlInfo:
    merges: List[Tuple[int, int, int, int]] = field(default_factory=list)
    hidden_rows: Set[int] = field(default_factory=set)
    hidden_cols: Set[int] = field(default_factory=set)
    uncached_formulas: Dict[Tuple[int, int], str] = field(default_factory=dict)


def workbook_sheet_paths(zf: zipfile.ZipFile) -> Dict[str, str]:
    """시트 이름 → 시트 XML 경로 (xl/worksheets/sheetN.xml)."""
    try:
        wb_xml = zf.read("xl/workbook.xml")
        rels_xml = zf.read("xl/_rels/workbook.xml.rels")
    except KeyError:
        return {}
    rels = {}
    for m in _TAG_REL.finditer(rels_xml):
        a = _attrs(m.group(1))
        if a.get("Id") and a.get("Target"):
            rels[a["Id"]] = a["Target"]
    out: Dict[str, str] = {}
    for m in _TAG_SHEET.finditer(wb_xml):
        a = _attrs(m.group(1))
        name, rid = a.get("name"), a.get("rid")
        target = rels.get(rid or "")
        if not name or not target:
            continue
        target = target.lstrip("/")
        out[name] = target if target.startswith("xl/") else f"xl/{target}"
    return out


def worksheets_xml_size(zf: zipfile.ZipFile) -> int:
    return sum(i.file_size for i in zf.infolist()
               if i.filename.startswith("xl/worksheets/") and i.filename.endswith(".xml"))


def scan_sheet_xml(data: bytes, want_structure: bool) -> SheetXmlInfo:
    """시트 XML 에서 계산값 없는 수식(늘), 병합·숨김 행열(want_structure 일 때)을 읽는다."""
    info = SheetXmlInfo()
    if want_structure:
        for m in _TAG_ROW.finditer(data):
            a = _attrs(m.group(1))
            if a.get("hidden") in ("1", "true") and a.get("r", "").isdigit():
                info.hidden_rows.add(int(a["r"]))
        for m in _TAG_COL.finditer(data):
            a = _attrs(m.group(1))
            if a.get("hidden") in ("1", "true"):
                try:
                    info.hidden_cols.update(range(int(a.get("min", "0")), int(a.get("max", "0")) + 1))
                except ValueError:
                    pass
        for m in _TAG_MERGE.finditer(data):
            ref = _attrs(m.group(1)).get("ref", "")
            try:
                a, _, b = ref.partition(":")
                r1, c1 = _parse_ref(a)
                r2, c2 = _parse_ref(b or a)
                info.merges.append((r1, c1, r2, c2))
            except ValueError:
                continue
    if b"<f" not in data and b":f" not in data:
        return info
    shared: Dict[str, Tuple[str, str]] = {}  # si -> (수식, 원래 셀)
    pending: List[Tuple[Tuple[int, int], str, str]] = []  # 공유 수식의 딸린 셀
    for m in _CELL.finditer(data):
        body = m.group(3)
        fm = _F.search(body)
        if not fm:
            continue
        a = _attrs(m.group(2))
        ref = a.get("r")
        if not ref:
            continue
        fa = _attrs(fm.group(1))
        ftext = html.unescape((fm.group(2) or b"").decode("utf-8", "replace"))
        if fa.get("t") == "shared" and fa.get("si") is not None and ftext:
            shared[fa["si"]] = (ftext, ref)
        vm = _V.search(body)
        has_value = vm is not None and vm.group(1).strip() != b""
        if a.get("t") == "inlineStr" or has_value:
            continue
        try:
            pos = _parse_ref(ref)
        except ValueError:
            continue
        if ftext:
            info.uncached_formulas[pos] = "=" + ftext
        elif fa.get("t") == "shared" and fa.get("si") is not None:
            pending.append((pos, ref, fa["si"]))
    if pending:
        try:
            from openpyxl.formula.translate import Translator
        except Exception:  # pragma: no cover
            Translator = None
        for pos, ref, si in pending:
            master = shared.get(si)
            if not master:
                continue
            ftext, origin = master
            if Translator is not None:
                try:
                    info.uncached_formulas[pos] = Translator("=" + ftext, origin=origin).translate_formula(ref)
                    continue
                except Exception:
                    pass
            info.uncached_formulas[pos] = "=" + ftext
    return info


# ──────────────────────────────────────────────────────────────────────────────
# xlsx 계열: openpyxl 워크시트 → SheetGrid
# ──────────────────────────────────────────────────────────────────────────────

_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA", "#SPILL!", "#CALC!"}


def _kind_of(value) -> str:
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        return "date"
    if isinstance(value, str) and value in _ERRORS:
        return "error"
    return "text"


def _has_border(cell) -> bool:
    try:
        b = cell.border
        if b is None:
            return False
        for side in (b.top, b.bottom, b.left, b.right):
            if side is not None and side.style not in (None, "none"):
                return True
    except Exception:
        return False
    return False


def grid_from_openpyxl(ws, xml_info: Optional[SheetXmlInfo] = None, *, read_only: bool = False,
                       date1904: bool = False) -> SheetGrid:
    """openpyxl 워크시트(전체 모드 또는 read_only) → SheetGrid."""
    grid = SheetGrid(name=ws.title, hidden=getattr(ws, "sheet_state", "visible") != "visible")
    uncached = xml_info.uncached_formulas if xml_info else {}
    if read_only:
        if xml_info:
            grid.merges = list(xml_info.merges)
            grid.hidden_rows = set(xml_info.hidden_rows)
            grid.hidden_cols = set(xml_info.hidden_cols)
    else:
        grid.merges = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges]
        for idx, dim in ws.row_dimensions.items():
            if getattr(dim, "hidden", False):
                grid.hidden_rows.add(idx)
        for dim in ws.column_dimensions.values():
            if getattr(dim, "hidden", False) and dim.min and dim.max:
                grid.hidden_cols.update(range(dim.min, dim.max + 1))

    check_borders = not read_only
    if check_borders:
        try:
            if (ws.max_row or 0) * (ws.max_column or 0) > BORDER_SCAN_MAX_CELLS:
                check_borders = False
        except Exception:
            check_borders = False

    cells = grid.cells
    for row in ws.iter_rows():
        for cell in row:
            r = getattr(cell, "row", None)
            c = getattr(cell, "column", None)
            if r is None or c is None:
                continue  # read_only 의 EmptyCell
            value = cell.value
            if value is None or (isinstance(value, str) and not value.strip()):
                formula = uncached.get((r, c))
                if formula:
                    cells[(r, c)] = GridCell(formula, "formula")
                elif check_borders and getattr(cell, "has_style", True) and _has_border(cell):
                    grid.bordered.add((r, c))
                continue
            kind = _kind_of(value)
            if kind == "text":
                text = _clean_text(value)
            else:
                try:
                    text = _clean_text(format_cell_value(value, cell.number_format, date1904=date1904))
                except Exception:
                    text = _clean_text(str(value))
            if text:
                cells[(r, c)] = GridCell(text, kind)
    for pos, formula in uncached.items():  # iter_rows 가 지나치지 않은 자리(드묾)
        cells.setdefault(pos, GridCell(formula, "formula"))
    return grid


# ──────────────────────────────────────────────────────────────────────────────
# xls: xlrd 시트 → SheetGrid
# ──────────────────────────────────────────────────────────────────────────────

def grid_from_xlrd(sheet, book) -> SheetGrid:
    import xlrd

    grid = SheetGrid(name=sheet.name, hidden=getattr(sheet, "visibility", 0) != 0)
    grid.merges = [(rlo + 1, clo + 1, rhi, chi) for (rlo, rhi, clo, chi) in sheet.merged_cells]
    try:
        for r, info in sheet.rowinfo_map.items():
            if getattr(info, "hidden", 0):
                grid.hidden_rows.add(r + 1)
        for c, info in sheet.colinfo_map.items():
            if getattr(info, "hidden", 0):
                grid.hidden_cols.add(c + 1)
    except Exception:
        pass

    formatting = bool(getattr(book, "formatting_info", False)) and bool(getattr(book, "xf_list", None))
    check_borders = formatting and sheet.nrows * sheet.ncols <= BORDER_SCAN_MAX_CELLS

    def fmt_of(r: int, c: int) -> Optional[str]:
        if not formatting:
            return None
        try:
            xf = book.xf_list[sheet.cell_xf_index(r, c)]
            return book.format_map[xf.format_key].format_str
        except Exception:
            return None

    def bordered(r: int, c: int) -> bool:
        try:
            b = book.xf_list[sheet.cell_xf_index(r, c)].border
            return any(s and s > 0 for s in (b.top_line_style, b.bottom_line_style,
                                              b.left_line_style, b.right_line_style))
        except Exception:
            return False

    for r in range(sheet.nrows):
        ctypes = sheet.row_types(r)
        values = sheet.row_values(r)
        for c, (ctype, value) in enumerate(zip(ctypes, values)):
            if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                if check_borders and ctype == xlrd.XL_CELL_BLANK and bordered(r, c):
                    grid.bordered.add((r + 1, c + 1))
                continue
            if ctype == xlrd.XL_CELL_TEXT:
                text, kind = str(value), "text"
            elif ctype == xlrd.XL_CELL_NUMBER:
                text, kind = format_cell_value(value, fmt_of(r, c)), "number"
            elif ctype == xlrd.XL_CELL_DATE:
                try:
                    text = format_cell_value(xlrd.xldate_as_datetime(value, book.datemode), fmt_of(r, c))
                except Exception:
                    text = format_cell_value(value, fmt_of(r, c), date1904=book.datemode == 1)
                kind = "date"
            elif ctype == xlrd.XL_CELL_BOOLEAN:
                text, kind = ("TRUE" if value else "FALSE"), "bool"
            elif ctype == xlrd.XL_CELL_ERROR:
                text, kind = xlrd.error_text_from_code.get(value, "#ERR"), "error"
            else:
                text, kind = str(value), "text"
            text = _clean_text(text)
            if text:
                grid.cells[(r + 1, c + 1)] = GridCell(text, kind)
    return grid


# ──────────────────────────────────────────────────────────────────────────────
# 워크북 열기
# ──────────────────────────────────────────────────────────────────────────────

def is_large_workbook(file_data: bytes, threshold: int = STREAMING_THRESHOLD_BYTES) -> bool:
    """시트 XML 합계가 threshold 를 넘는가(넘으면 read_only 스트리밍으로 읽는다)."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_data)) as zf:
            return worksheets_xml_size(zf) > threshold
    except zipfile.BadZipFile:
        return False


def load_streaming_workbook(file_data: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(file_data), read_only=True, data_only=True, keep_links=False)


def grids_from_workbook(wb, file_data: bytes, *, read_only: bool) -> List[SheetGrid]:
    """열린 openpyxl 워크북의 시트마다 SheetGrid. file_data 는 시트 XML 을 직접 읽는 데 쓴다."""
    date1904 = getattr(getattr(wb, "epoch", None), "year", 1900) == 1904
    try:
        zf = zipfile.ZipFile(io.BytesIO(file_data))
        paths = workbook_sheet_paths(zf)
    except zipfile.BadZipFile:
        zf, paths = None, {}
    grids: List[SheetGrid] = []
    for ws in wb.worksheets:
        info = None
        path = paths.get(ws.title)
        if zf is not None and path:
            try:
                info = scan_sheet_xml(zf.read(path), want_structure=read_only)
            except KeyError:
                info = None
        grids.append(grid_from_openpyxl(ws, info, read_only=read_only, date1904=date1904))
    return grids


def xls_grids(book) -> List[SheetGrid]:
    return [grid_from_xlrd(book.sheet_by_index(i), book) for i in range(book.nsheets)]


# ──────────────────────────────────────────────────────────────────────────────
# 엑셀 확장자로 내보낸 HTML 표(관공서 사이트 등)
# ──────────────────────────────────────────────────────────────────────────────

_NUMBER_TEXT = re.compile(r"^[-+]?(?:\d{1,3}(?:,\d{3})+|\d+)?(?:\.\d+)?%?$")
_DATE_TEXT = re.compile(r"^\d{4}[-./년]\s?\d{1,2}[-./월]\s?\d{1,2}")


def _guess_kind(text: str) -> str:
    """글자만 있는 셀(HTML·CSV)의 값 종류를 모양으로 짐작한다."""
    if any(ch.isdigit() for ch in text) and _NUMBER_TEXT.match(text):
        return "number"
    if _DATE_TEXT.match(text):
        return "date"
    return "text"


def _span(cell, attr: str) -> int:
    try:
        return max(1, min(int(str(cell.get(attr, 1)).strip() or 1), 1000))
    except (TypeError, ValueError):
        return 1


def grid_from_html_rows(name: str, rows) -> SheetGrid:
    """BeautifulSoup ``<tr>`` 목록 → SheetGrid. colspan·rowspan 은 병합으로 남긴다
    (값을 병합 칸마다 복사하지 않는다. 머리글 평탄화와 세로 병합 채움은 렌더러가 한다)."""
    grid = SheetGrid(name=name)
    taken: Set[Tuple[int, int]] = set()
    for r, tr in enumerate(rows, start=1):
        c = 1
        for cell in tr.find_all(["td", "th"]):
            while (r, c) in taken:
                c += 1
            rs, cs = _span(cell, "rowspan"), _span(cell, "colspan")
            text = _clean_text(cell.get_text(" ", strip=True))
            if text:
                grid.cells[(r, c)] = GridCell(text, _guess_kind(text))
            for dr in range(rs):
                for dc in range(cs):
                    taken.add((r + dr, c + dc))
            if rs > 1 or cs > 1:
                grid.merges.append((r, c, r + rs - 1, c + cs - 1))
            c += cs
    return grid
