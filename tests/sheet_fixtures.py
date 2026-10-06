"""
Spreadsheet fixtures built in code (no binary files in the repository).

Every builder takes a directory and returns the path of the file it wrote.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Side

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NAMES = ["김민준", "이서연", "박지호", "최수아", "정우진", "강하은", "조예준", "윤지유"]


def _save(wb: Workbook, path: Path) -> Path:
    wb.save(path)
    return path


def many_rows_xlsx(d: Path, rows: int = 3000) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "주문"
    ws.append(["주문번호", "고객", "수량", "금액"])
    for i in range(1, rows + 1):
        ws.append([f"ORD-{i:05d}", NAMES[i % len(NAMES)], i % 7 + 1, i * 1000])
    return _save(wb, d / "many_rows.xlsx")


def many_rows_xls(d: Path, rows: int = 1500) -> Path:
    import xlwt

    book = xlwt.Workbook()
    sh = book.add_sheet("주문")
    for c, h in enumerate(["주문번호", "고객", "수량", "금액"]):
        sh.write(0, c, h)
    for i in range(1, rows + 1):
        sh.write(i, 0, f"ORD-{i:05d}")
        sh.write(i, 1, NAMES[i % len(NAMES)])
        sh.write(i, 2, i % 7 + 1)
        sh.write(i, 3, i * 1000)
    path = d / "many_rows.xls"
    book.save(str(path))
    return path


def wide_xlsx(d: Path, cols: int = 130) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "센서"
    ws.append([f"S{c:03d}" for c in range(1, cols + 1)])
    for r in range(1, 6):
        ws.append([r * 1000 + c for c in range(1, cols + 1)])
    return _save(wb, d / "wide.xlsx")


def merged_header_xlsx(d: Path, rows: int = 120) -> Path:
    """Two-row header: 지점 | 담당 | 1월(목표, 실적) | 2월(목표, 실적)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "실적"
    ws["A1"], ws["B1"], ws["C1"], ws["E1"] = "지점", "담당", "1월", "2월"
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:D1")
    ws.merge_cells("E1:F1")
    ws["C2"], ws["D2"], ws["E2"], ws["F2"] = "목표", "실적", "목표", "실적"
    for i in range(rows):
        ws.append([f"지점{i + 1:03d}", NAMES[i % len(NAMES)], 100 + i, 90 + i, 110 + i, 95 + i])
    return _save(wb, d / "merged_header.xlsx")


def gap_column_xlsx(d: Path) -> Path:
    """A table with an empty spacer column (C) between B and D."""
    wb = Workbook()
    ws = wb.active
    ws.title = "재고"
    ws.append(["품목", "수량", None, "창고", "비고"])
    for i in range(10):
        ws.append([f"품목{i}", i * 3, None, f"창고{i % 3}", "정상" if i % 2 else "점검"])
    return _save(wb, d / "gap_column.xlsx")


def blank_row_groups_xlsx(d: Path) -> Path:
    """One table whose groups are separated by blank rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "부서"
    ws.append(["부서", "이름", "직급"])
    r = 2
    for g, dept in enumerate(["영업", "개발", "인사"]):
        for k in range(3):
            ws.cell(r, 1, dept)
            ws.cell(r, 2, NAMES[(g * 3 + k) % len(NAMES)])
            ws.cell(r, 3, "팀장" if k == 0 else "사원")
            r += 1
        r += 1  # blank row between groups
    return _save(wb, d / "blank_row_groups.xlsx")


def form_xlsx(d: Path) -> Path:
    """Form layout: label | spacer | value."""
    wb = Workbook()
    ws = wb.active
    ws.title = "신청서"
    ws["A1"] = "출장 신청서"
    ws.merge_cells("A1:C1")
    fields = [("신청자", "김민준"), ("부서", "영업1팀"), ("출장지", "부산"), ("기간", "2026-10-07 ~ 2026-10-09")]
    for i, (k, v) in enumerate(fields, start=3):
        ws.cell(i, 1, k)
        ws.cell(i, 3, v)
    return _save(wb, d / "form.xlsx")


def sparse_xlsx(d: Path) -> Path:
    """Attendance matrix with sparse "O" marks and bordered empty cells."""
    wb = Workbook()
    ws = wb.active
    ws.title = "출석"
    days = ["월", "화", "수", "목", "금"]
    ws.append(["이름"] + days)
    for i, name in enumerate(NAMES[:5]):
        ws.append([name] + ["O" if (i + j) % 3 == 0 else None for j in range(len(days))])
    for row in ws.iter_rows(min_row=1, max_row=6, min_col=1, max_col=6):
        for cell in row:
            cell.border = BOX
    return _save(wb, d / "sparse.xlsx")


def formats_xlsx(d: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "서식"
    ws.append(["항목", "값"])
    rows = [
        ("비율", 0.153, "0.0%"),
        ("금액", 1234567, '#,##0"원"'),
        ("손익", -1500, "#,##0;(#,##0)"),
        ("일자", dt.datetime(2026, 10, 6), 'yyyy"년" m"월" d"일"'),
        ("요일", dt.datetime(2026, 10, 6), "yyyy-mm-dd (aaa)"),
        ("시각", dt.time(14, 5), "h:mm AM/PM"),
        ("통화", 1234.5, "[$₩-412]#,##0.00"),
        ("천원", 12345678, '#,##0,"천원"'),
        ("코드", 42, "00000"),
    ]
    for label, value, fmt in rows:
        ws.append([label, value])
        ws.cell(ws.max_row, 2).number_format = fmt
    return _save(wb, d / "formats.xlsx")


def formulas_xlsx(d: Path, cached: bool = False) -> Path:
    """Formulas saved by openpyxl have no cached values. With cached=True the
    sheet XML gets <v> values injected, as Excel would write them."""
    wb = Workbook()
    ws = wb.active
    ws.title = "계산"
    ws.append(["품목", "단가", "수량", "합계"])
    for i in range(1, 4):
        ws.append([f"품목{i}", i * 100, i, f"=B{i + 1}*C{i + 1}"])
    path = d / ("formulas_cached.xlsx" if cached else "formulas.xlsx")
    wb.save(path)
    if cached:
        _inject_cached_values(path, {f"D{i + 1}": i * 100 * i for i in range(1, 4)})
    return path


def _inject_cached_values(path: Path, values: dict) -> None:
    src = path.read_bytes()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                text = data.decode("utf-8")
                for ref, v in values.items():
                    text = re.sub(
                        rf'(<c r="{ref}"[^>]*>)(<f>[^<]*</f>)(<v></v>|<v/>)?',
                        lambda m: f"{m.group(1)}{m.group(2)}<v>{v}</v>",
                        text,
                    )
                data = text.encode("utf-8")
            zout.writestr(item, data)
    path.write_bytes(out.getvalue())


def hidden_xlsx(d: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "공개"
    ws.append(["이름", "연락처", "메모"])
    ws.append(["김민준", "010-0000-0001", "보이는 행"])
    ws.append(["이서연", "010-0000-0002", "숨긴 행"])
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["C"].hidden = True
    hid = wb.create_sheet("숨김시트")
    hid.append(["비밀", "값"])
    hid.append(["내부", 1])
    hid.sheet_state = "hidden"
    return _save(wb, d / "hidden.xlsx")


def xlsm(d: Path) -> Path:
    """A macro-enabled workbook (content type of an .xlsm, no VBA part)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "매크로"
    ws.append(["코드", "설명"])
    ws.append(["M-01", "월말 정산"])
    buf = io.BytesIO()
    wb.save(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            zout.writestr(item, data)
    path = d / "macro.xlsm"
    path.write_bytes(out.getvalue())
    return path


def multi_table_xlsx(d: Path) -> Path:
    """Two captioned tables on one sheet, stacked with blank rows between."""
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = "2026년 3분기 매출"
    ws.merge_cells("A1:C1")
    ws.append(["지역", "매출", "비중"])
    for i, region in enumerate(["서울", "부산", "대구"]):
        ws.append([region, 1000 * (i + 1), f"{(i + 1) * 10}%"])
    ws.append([])
    ws.append([])
    ws.append(["2026년 3분기 비용"])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=3)
    ws.append(["항목", "금액", "비고"])
    for item in ["인건비", "임차료", "광고비"]:
        ws.append([item, 500, "-"])
    return _save(wb, d / "multi_table.xlsx")


def long_cell_xlsx(d: Path, length: int = 5000) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "설명"
    ws.append(["번호", "제목", "본문"])
    sentence = "이 문서는 긴 셀 처리를 확인하기 위한 문장입니다. "
    body = (sentence * (length // len(sentence) + 1))[:length]
    ws.append([1, "긴 본문", body])
    ws.append([2, "짧은 본문", "짧다"])
    return _save(wb, d / "long_cell.xlsx")


def image_xlsx(d: Path) -> Path:
    """Sheet 1 holds an image, sheet 2 has none."""
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image

    png = d / "logo.png"
    Image.new("RGB", (40, 20), (200, 30, 30)).save(png)
    wb = Workbook()
    ws = wb.active
    ws.title = "로고"
    ws.append(["회사", "로고"])
    ws.append(["플래티어", None])
    ws.add_image(XLImage(str(png)), "B2")
    ws2 = wb.create_sheet("목록")
    ws2.append(["번호", "이름"])
    ws2.append([1, "가"])
    return _save(wb, d / "image.xlsx")


def xls_merged(d: Path) -> Path:
    import xlwt

    book = xlwt.Workbook()
    sh = book.add_sheet("실적")
    sh.write_merge(0, 1, 0, 0, "지점")
    sh.write_merge(0, 0, 1, 2, "1월")
    sh.write(1, 1, "목표")
    sh.write(1, 2, "실적")
    pct = xlwt.easyxf(num_format_str="0.0%")
    date = xlwt.easyxf(num_format_str="yyyy-mm-dd")
    for i in range(5):
        sh.write(i + 2, 0, f"지점{i + 1}")
        sh.write(i + 2, 1, 0.5 + i / 100, pct)
        sh.write(i + 2, 2, dt.datetime(2026, 1, i + 1), date)
    path = d / "merged.xls"
    book.save(str(path))
    return path


def csv_blank_cells(d: Path) -> Path:
    """Employee list where some 사번 cells are empty (must not borrow the row above)."""
    text = "사번,이름,부서\n1001,김민준,영업\n1002,이서연,개발\n,최수아,개발\n1003,박지호,\n"
    path = d / "employees.csv"
    path.write_text(text, encoding="utf-8")
    return path


def csv_many_rows(d: Path, rows: int = 3000) -> Path:
    lines = ["번호,이름,점수"] + [f"{i},{NAMES[i % len(NAMES)]},{i % 100}" for i in range(1, rows + 1)]
    path = d / "many_rows.csv"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def tsv(d: Path) -> Path:
    path = d / "items.tsv"
    path.write_text("코드\t이름\n A1\t사과\nB2\t배\n", encoding="utf-8")
    return path


def html_xls(d: Path) -> Path:
    """An HTML table saved with an .xls extension (government site exports):
    a header-only table followed by a body-only table, with a two-row merged header."""
    rows = "".join(
        f"<tr><td>{2020 + i}</td><td>{100 + i}</td><td>{90 + i}</td><td>{i}%</td></tr>" for i in range(5)
    )
    html_text = (
        "<html><head><meta charset='utf-8'></head><body>"
        "<table><thead>"
        "<tr><th rowspan='2'>연도</th><th colspan='2'>인구(천 명)</th><th rowspan='2'>증감률</th></tr>"
        "<tr><th>남</th><th>여</th></tr>"
        "</thead></table>"
        f"<table><tbody>{rows}</tbody></table>"
        "</body></html>"
    )
    path = d / "통계.xls"
    path.write_text(html_text, encoding="utf-8")
    return path
