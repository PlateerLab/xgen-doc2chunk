"""End-to-end checks for spreadsheets: extract_text -> chunk_text as xgen-documents calls them."""
import re

import pytest

import sheet_fixtures as F
from conftest import extract


def _rows(chunks, pattern):
    return set(re.findall(pattern, "".join(c["text"] for c in chunks)))


@pytest.mark.parametrize("builder, count", [(F.many_rows_xlsx, 3000), (F.many_rows_xls, 1500)])
def test_every_row_is_kept(processor, tmp_path, builder, count):
    _, chunks = extract(processor, builder(tmp_path))
    assert len(_rows(chunks, r"ORD-\d{5}")) == count


def test_every_column_is_kept(processor, tmp_path):
    text, _ = extract(processor, F.wide_xlsx(tmp_path))
    assert "S130" in text and "1130" in text


def test_table_chunks_respect_chunk_size(processor, tmp_path):
    _, chunks = extract(processor, F.many_rows_xlsx(tmp_path))
    assert max(len(c["text"]) for c in chunks) <= 1000


def test_merged_header_in_every_chunk(processor, tmp_path):
    _, chunks = extract(processor, F.merged_header_xlsx(tmp_path))
    assert len(chunks) > 1
    header = "| 지점 | 담당 | 1월 목표 | 1월 실적 | 2월 목표 | 2월 실적 |"
    assert all(header in c["text"] for c in chunks)


def test_rows_have_full_width(processor, tmp_path):
    """Pipe-table parsers (xgen-documents ontology) keep only rows of the header width."""
    _, chunks = extract(processor, F.sparse_xlsx(tmp_path))
    rows = [l for c in chunks for l in c["text"].splitlines() if l.startswith("|")]
    assert len({l.count("|") for l in rows}) == 1
    assert "| 이서연 |  |  | O |  |  |" in rows


def test_spacer_column_does_not_split_table(processor, tmp_path):
    text, _ = extract(processor, F.gap_column_xlsx(tmp_path))
    assert "| 품목 | 수량 | 창고 | 비고 |" in text
    assert "| 품목3 | 9 | 창고0 | 정상 |" in text


def test_blank_rows_do_not_split_table(processor, tmp_path):
    text, _ = extract(processor, F.blank_row_groups_xlsx(tmp_path))
    assert text.count("| 부서 | 이름 | 직급 |") == 1
    assert "| 인사 | 조예준 | 팀장 |" in text


def test_form_layout_becomes_label_value(processor, tmp_path):
    text, _ = extract(processor, F.form_xlsx(tmp_path))
    assert "출장 신청서" in text
    assert "신청자: 김민준" in text
    assert "기간: 2026-10-07 ~ 2026-10-09" in text


def test_display_formats(processor, tmp_path):
    text, _ = extract(processor, F.formats_xlsx(tmp_path))
    for shown in ["15.3%", "1,234,567원", "(1,500)", "2026년 10월 6일", "2026-10-06 (화)",
                  "2:05 PM", "₩1,234.50", "12,346천원", "00042"]:
        assert f"| {shown} |" in text, shown


def test_formula_without_cached_value_is_kept(processor, tmp_path):
    text, _ = extract(processor, F.formulas_xlsx(tmp_path))
    assert "| 품목2 | 200 | 2 | =B3*C3 |" in text


def test_formula_cached_value_wins(processor, tmp_path):
    text, _ = extract(processor, F.formulas_xlsx(tmp_path, cached=True))
    assert "| 품목2 | 200 | 2 | 400 |" in text


def test_macro_workbook(processor, tmp_path):
    text, _ = extract(processor, F.xlsm(tmp_path))
    assert "| M-01 | 월말 정산 |" in text


def test_captions_and_multiple_tables(processor, tmp_path):
    text, _ = extract(processor, F.multi_table_xlsx(tmp_path))
    assert "[Table 1] 2026년 3분기 매출\n| 지역 | 매출 | 비중 |" in text
    assert "[Table 2] 2026년 3분기 비용\n| 항목 | 금액 | 비고 |" in text


def test_caption_repeated_when_table_is_split(processor, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "거래처 목록"
    ws.merge_cells("A1:C1")
    ws.append(["코드", "상호", "지역"])
    for i in range(200):
        ws.append([f"C{i:04d}", f"상호{i}", "서울"])
    path = tmp_path / "caption.xlsx"
    wb.save(path)
    _, chunks = extract(processor, path)
    assert len(chunks) > 1
    assert all("[Table 1] 거래처 목록" in c["text"] and "| 코드 | 상호 | 지역 |" in c["text"] for c in chunks)


def test_long_cell_row_is_split_with_other_cells_repeated(processor, tmp_path):
    _, chunks = extract(processor, F.long_cell_xlsx(tmp_path))
    pieces = [l for c in chunks for l in c["text"].splitlines() if l.startswith("| 1 | 긴 본문 |")]
    assert len(pieces) > 1
    assert max(len(c["text"]) for c in chunks) <= 1000


def test_hidden_content_included_by_default(processor, tmp_path):
    text, _ = extract(processor, F.hidden_xlsx(tmp_path))
    assert "숨긴 행" in text and "[Sheet: 숨김시트]" in text


@pytest.mark.parametrize("streaming", [False, True])
def test_hidden_content_can_be_excluded(make_processor, tmp_path, streaming):
    config = {"spreadsheet": {"include_hidden": False}}
    if streaming:
        config["spreadsheet"]["streaming_threshold_bytes"] = 1
    text, _ = extract(make_processor(config), F.hidden_xlsx(tmp_path))
    assert "김민준" in text
    assert "이서연" not in text      # hidden row
    assert "메모" not in text         # hidden column
    assert "숨김시트" not in text     # hidden sheet


def test_streaming_mode_matches_full_mode(make_processor, tmp_path):
    path = F.merged_header_xlsx(tmp_path)
    full, _ = extract(make_processor(), path)
    streamed, _ = extract(make_processor({"spreadsheet": {"streaming_threshold_bytes": 1}}), path)
    assert streamed == full
    path = F.formats_xlsx(tmp_path)
    assert extract(make_processor({"spreadsheet": {"streaming_threshold_bytes": 1}}), path)[0] == \
        extract(make_processor(), path)[0]


def test_sheet_images_are_not_repeated(processor, tmp_path):
    text, chunks = extract(processor, F.image_xlsx(tmp_path))
    assert text.count("[Image:") == 1
    first = next(c for c in chunks if "[Sheet: 로고]" in c["text"])
    assert "[Image:" in first["text"]


def test_sheet_page_numbers(processor, tmp_path):
    _, chunks = extract(processor, F.hidden_xlsx(tmp_path))
    assert [c["page_number"] for c in chunks] == [1, 2]


def test_xls_merged_header_and_formats(processor, tmp_path):
    text, _ = extract(processor, F.xls_merged(tmp_path))
    assert "| 지점 | 1월 목표 | 1월 실적 |" in text
    assert "| 지점2 | 51.0% | 2026-01-02 |" in text


def test_pipe_in_cell_is_escaped(processor, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["키", "값"])
    ws.append(["경로", "a|b"])
    path = tmp_path / "pipe.xlsx"
    wb.save(path)
    text, _ = extract(processor, path)
    assert "| 경로 | a\\|b |" in text


def test_html_saved_as_xls(processor, tmp_path):
    text, chunks = extract(processor, F.html_xls(tmp_path))
    assert "<table" not in text
    assert "[Sheet: 통계]" in text
    assert "| 연도 | 인구(천 명) 남 | 인구(천 명) 여 | 증감률 |" in text
    assert "| 2022 | 102 | 92 | 2% |" in text
